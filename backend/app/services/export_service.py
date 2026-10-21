from io import BytesIO, StringIO
import csv
from typing import List, Tuple
from datetime import timezone, timedelta
from app.models.order import Order

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

# IST Timezone (UTC + 5:30)
IST = timezone(timedelta(hours=5, minutes=30))

class ExportService:
    @staticmethod
    def generate_orders_excel(orders: List[Order]) -> Tuple[BytesIO, str, str]:
        """
        Returns (buffer, media_type, file_extension)
        Uses openpyxl (.xlsx) if available, or UTF-8 BOM CSV (.csv) as fallback.
        """
        if not HAS_OPENPYXL:
            return ExportService._generate_orders_csv(orders)
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Orders & Dispatch"
        ws.views.sheetView[0].showGridLines = True

        # Styles
        header_fill = PatternFill(start_color="181821", end_color="181821", fill_type="solid")
        header_font = Font(name="Segoe UI", size=10, bold=True, color="D4AF37")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        data_font = Font(name="Segoe UI", size=9, color="000000")
        bold_data_font = Font(name="Segoe UI", size=9, bold=True, color="000000")
        engraving_font = Font(name="Segoe UI", size=9, bold=True, color="B8860B")
        
        thin_border = Border(
            left=Side(style="thin", color="E0E0E0"),
            right=Side(style="thin", color="E0E0E0"),
            top=Side(style="thin", color="E0E0E0"),
            bottom=Side(style="thin", color="E0E0E0")
        )

        header_border = Border(
            left=Side(style="thin", color="333333"),
            right=Side(style="thin", color="333333"),
            top=Side(style="medium", color="D4AF37"),
            bottom=Side(style="medium", color="D4AF37")
        )

        headers = [
            "Order ID",
            "Order Date & Time (IST)",
            "Order Status",
            "Payment Method",
            "Payment Status",
            "Customer Name",
            "Customer Phone",
            "Customer Email",
            "Address Line 1",
            "Address Line 2",
            "City",
            "State",
            "Pincode",
            "Full Address",
            "Product Name",
            "SKU",
            "Qty",
            "Unit Price (Rs)",
            "Item Total (Rs)",
            "Bat Weight",
            "Handle Shape",
            "Handle Size",
            "Grip Color",
            "Grip Pattern",
            "Grip Count",
            "Sticker Finish",
            "Pre-Knocking",
            "Oiling",
            "Face Protection",
            "Custom Engraving",
            "All Customization Specs",
            "Order Subtotal (Rs)",
            "GST Rate (%)",
            "GST Amount (Rs)",
            "Shipping Fee (Rs)",
            "Coupon Code",
            "Discount (Rs)",
            "Grand Total (Rs)",
            "Tracking Number",
            "Shipping Carrier",
            "Customer Notes",
            "Admin Notes"
        ]

        ws.append(headers)
        ws.row_dimensions[1].height = 30

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = header_border

        current_row = 2
        total_items_count = 0
        total_revenue = 0.0

        for order in orders:
            # Format order date into IST
            created_dt = order.created_at
            if created_dt:
                if created_dt.tzinfo is None:
                    # Treat naive as UTC then convert
                    created_dt = created_dt.replace(tzinfo=timezone.utc).astimezone(IST)
                else:
                    created_dt = created_dt.astimezone(IST)
                date_str = created_dt.strftime("%Y-%m-%d %H:%M:%S")
            else:
                date_str = ""

            shipping = order.shipping_address or {}
            addr_line1 = shipping.get("address_line1", "") or ""
            addr_line2 = shipping.get("address_line2", "") or ""
            city = shipping.get("city", "") or ""
            state = shipping.get("state", "") or ""
            pincode = str(shipping.get("pincode", "") or "")

            address_parts = [addr_line1]
            if addr_line2:
                address_parts.append(addr_line2)
            if city:
                address_parts.append(city)
            if state:
                address_parts.append(state)
            if pincode:
                address_parts.append(pincode)
            full_address_str = ", ".join(address_parts)

            subtotal_val = float(order.subtotal or 0.0)
            gst_percent_val = float(order.gst_percent or 0.0)
            gst_amount_val = float(order.gst_amount or 0.0)
            shipping_val = float(order.shipping_fee or 0.0)
            discount_val = float(order.discount_amount or 0.0)
            grand_total_val = float(order.grand_total or 0.0)

            total_revenue += grand_total_val

            # Handle orders with items
            items = order.items if order.items else [None]
            for item in items:
                total_items_count += (item.quantity if item else 1)

                prod_name = item.product_name if item else "N/A"
                sku = item.product_sku if item else "N/A"
                qty = item.quantity if item else 1
                unit_price = float(item.unit_price or 0.0) if item else 0.0
                item_total = float(item.total_price or 0.0) if item else 0.0

                custom = (item.customization or {}) if item else {}
                weight = custom.get("weight", "") or ""
                handle_shape = custom.get("handle_shape", "") or ""
                handle_size = custom.get("handle_size", "") or ""
                grip_color = custom.get("grip_color", "") or ""
                grip_pattern = custom.get("grip_pattern", "") or ""
                grip_count = custom.get("grip_count", "") or ""
                sticker_finish = custom.get("sticker_finish", "") or ""
                pre_knocking = custom.get("pre_knocking", "") or ""
                oiling = custom.get("oiling", "") or ""
                face_protection = custom.get("face_protection", "") or ""
                engraving = custom.get("custom_engraving", "") or ""

                specs_list = []
                if weight: specs_list.append(f"Weight: {weight}")
                if handle_shape: specs_list.append(f"Handle: {handle_shape} ({handle_size or 'SH'})")
                if grip_color: specs_list.append(f"Grip: {grip_color} {grip_pattern} ({grip_count})")
                if sticker_finish: specs_list.append(f"Sticker: {sticker_finish}")
                if pre_knocking and pre_knocking != "None": specs_list.append(f"Knocking: {pre_knocking}")
                if oiling and oiling != "None": specs_list.append(f"Oiling: {oiling}")
                if face_protection and face_protection != "None": specs_list.append(f"Sheet: {face_protection}")
                if engraving: specs_list.append(f"Engraving: \"{engraving}\"")
                specs_summary = " | ".join(specs_list) if specs_list else "Standard Specifications"

                row_values = [
                    order.order_number,
                    date_str,
                    (order.order_status or "pending").upper(),
                    (order.payment_method or "cod").upper(),
                    (order.payment_status or "pending").upper(),
                    shipping.get("full_name") or order.customer_name or "",
                    shipping.get("phone") or order.customer_phone or "",
                    shipping.get("email") or order.customer_email or "",
                    addr_line1,
                    addr_line2,
                    city,
                    state,
                    pincode,
                    full_address_str,
                    prod_name,
                    sku,
                    qty,
                    unit_price,
                    item_total,
                    weight,
                    handle_shape,
                    handle_size,
                    grip_color,
                    grip_pattern,
                    grip_count,
                    sticker_finish,
                    pre_knocking,
                    oiling,
                    face_protection,
                    engraving,
                    specs_summary,
                    subtotal_val,
                    gst_percent_val,
                    gst_amount_val,
                    shipping_val,
                    order.coupon_code or "None",
                    discount_val,
                    grand_total_val,
                    order.tracking_number or "Unassigned",
                    order.shipping_carrier or "Unassigned",
                    order.customer_notes or "",
                    order.admin_notes or ""
                ]

                ws.append(row_values)
                ws.row_dimensions[current_row].height = 20

                # Row styling
                is_even = (current_row % 2 == 0)
                row_fill = PatternFill(start_color="FAFAFA" if is_even else "FFFFFF", end_color="FAFAFA" if is_even else "FFFFFF", fill_type="solid")

                for c_idx in range(1, len(row_values) + 1):
                    c = ws.cell(row=current_row, column=c_idx)
                    c.font = data_font
                    c.fill = row_fill
                    c.border = thin_border

                    # Center align IDs, dates, statuses, pincode, etc.
                    if c_idx in [1, 2, 3, 4, 5, 13, 16, 17, 33, 36]:
                        c.alignment = Alignment(horizontal="center", vertical="center")
                    # Right align monetary values
                    elif c_idx in [18, 19, 32, 34, 35, 37, 38]:
                        c.alignment = Alignment(horizontal="right", vertical="center")
                        c.number_format = "#,##0.00"
                    else:
                        c.alignment = Alignment(horizontal="left", vertical="center")

                    # Highlight engraving if present
                    if c_idx == 30 and engraving:
                        c.font = engraving_font

                    # Bold order number and grand total
                    if c_idx in [1, 38]:
                        c.font = bold_data_font

                current_row += 1

        # Summary Row at bottom
        summary_row = current_row
        ws.row_dimensions[summary_row].height = 24
        summary_fill = PatternFill(start_color="F0E68C", end_color="F0E68C", fill_type="solid")
        summary_font = Font(name="Segoe UI", size=10, bold=True, color="000000")
        double_top_border = Border(
            top=Side(style="double", color="333333"),
            bottom=Side(style="medium", color="333333")
        )

        ws.cell(row=summary_row, column=1, value="TOTAL ORDERS:")
        ws.cell(row=summary_row, column=2, value=len(orders))
        ws.cell(row=summary_row, column=16, value="TOTAL ITEMS:")
        ws.cell(row=summary_row, column=17, value=total_items_count)
        ws.cell(row=summary_row, column=37, value="TOTAL REVENUE:")
        revenue_cell = ws.cell(row=summary_row, column=38, value=total_revenue)
        revenue_cell.number_format = "#,##0.00"

        for c_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=summary_row, column=c_idx)
            cell.font = summary_font
            cell.fill = summary_fill
            cell.border = double_top_border

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                # Ignore very long summary text for column width calculation to prevent huge columns
                if len(val_str) > max_len and len(val_str) < 60:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # Specifically set generous width for important columns
        ws.column_dimensions["A"].width = 22  # Order ID
        ws.column_dimensions["B"].width = 22  # Date
        ws.column_dimensions["F"].width = 22  # Customer Name
        ws.column_dimensions["G"].width = 16  # Phone
        ws.column_dimensions["H"].width = 24  # Email
        ws.column_dimensions["N"].width = 38  # Full Address
        ws.column_dimensions["O"].width = 26  # Product Name
        ws.column_dimensions["T"].width = 16  # Bat Weight
        ws.column_dimensions["AD"].width = 20 # Custom Engraving
        ws.column_dimensions["AE"].width = 45 # All Customization Specs
        ws.column_dimensions["AL"].width = 24 # Grand Total

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"

    @staticmethod
    def _generate_orders_csv(orders: List[Order]) -> Tuple[BytesIO, str, str]:
        output = StringIO()
        writer = csv.writer(output)

        headers = [
            "Order ID", "Order Date & Time (IST)", "Order Status", "Payment Method", "Payment Status",
            "Customer Name", "Customer Phone", "Customer Email", "Address Line 1", "Address Line 2",
            "City", "State", "Pincode", "Full Address", "Product Name", "SKU", "Qty",
            "Unit Price (Rs)", "Item Total (Rs)", "Bat Weight", "Handle Shape", "Handle Size",
            "Grip Color", "Grip Pattern", "Grip Count", "Sticker Finish", "Pre-Knocking",
            "Oiling", "Face Protection", "Custom Engraving", "All Customization Specs",
            "Order Subtotal (Rs)", "GST Rate (%)", "GST Amount (Rs)", "Shipping Fee (Rs)",
            "Coupon Code", "Discount (Rs)", "Grand Total (Rs)", "Tracking Number",
            "Shipping Carrier", "Customer Notes", "Admin Notes"
        ]
        writer.writerow(headers)

        for order in orders:
            created_dt = order.created_at
            if created_dt:
                if created_dt.tzinfo is None:
                    created_dt = created_dt.replace(tzinfo=timezone.utc).astimezone(IST)
                else:
                    created_dt = created_dt.astimezone(IST)
                date_str = created_dt.strftime("%Y-%m-%d %H:%M:%S")
            else:
                date_str = ""

            shipping = order.shipping_address or {}
            addr_line1 = shipping.get("address_line1", "") or ""
            addr_line2 = shipping.get("address_line2", "") or ""
            city = shipping.get("city", "") or ""
            state = shipping.get("state", "") or ""
            pincode = str(shipping.get("pincode", "") or "")
            address_parts = [addr_line1]
            if addr_line2: address_parts.append(addr_line2)
            if city: address_parts.append(city)
            if state: address_parts.append(state)
            if pincode: address_parts.append(pincode)
            full_address_str = ", ".join(address_parts)

            subtotal_val = float(order.subtotal or 0.0)
            gst_percent_val = float(order.gst_percent or 0.0)
            gst_amount_val = float(order.gst_amount or 0.0)
            shipping_val = float(order.shipping_fee or 0.0)
            discount_val = float(order.discount_amount or 0.0)
            grand_total_val = float(order.grand_total or 0.0)

            items = order.items if order.items else [None]
            for item in items:
                prod_name = item.product_name if item else "N/A"
                sku = item.product_sku if item else "N/A"
                qty = item.quantity if item else 1
                unit_price = float(item.unit_price or 0.0) if item else 0.0
                item_total = float(item.total_price or 0.0) if item else 0.0

                custom = (item.customization or {}) if item else {}
                weight = custom.get("weight", "") or ""
                handle_shape = custom.get("handle_shape", "") or ""
                handle_size = custom.get("handle_size", "") or ""
                grip_color = custom.get("grip_color", "") or ""
                grip_pattern = custom.get("grip_pattern", "") or ""
                grip_count = custom.get("grip_count", "") or ""
                sticker_finish = custom.get("sticker_finish", "") or ""
                pre_knocking = custom.get("pre_knocking", "") or ""
                oiling = custom.get("oiling", "") or ""
                face_protection = custom.get("face_protection", "") or ""
                engraving = custom.get("custom_engraving", "") or ""

                specs_list = []
                if weight: specs_list.append(f"Weight: {weight}")
                if handle_shape: specs_list.append(f"Handle: {handle_shape} ({handle_size or 'SH'})")
                if grip_color: specs_list.append(f"Grip: {grip_color} {grip_pattern} ({grip_count})")
                if sticker_finish: specs_list.append(f"Sticker: {sticker_finish}")
                if pre_knocking and pre_knocking != "None": specs_list.append(f"Knocking: {pre_knocking}")
                if oiling and oiling != "None": specs_list.append(f"Oiling: {oiling}")
                if face_protection and face_protection != "None": specs_list.append(f"Sheet: {face_protection}")
                if engraving: specs_list.append(f"Engraving: \"{engraving}\"")
                specs_summary = " | ".join(specs_list) if specs_list else "Standard Specifications"

                writer.writerow([
                    order.order_number,
                    date_str,
                    (order.order_status or "pending").upper(),
                    (order.payment_method or "cod").upper(),
                    (order.payment_status or "pending").upper(),
                    shipping.get("full_name") or order.customer_name or "",
                    shipping.get("phone") or order.customer_phone or "",
                    shipping.get("email") or order.customer_email or "",
                    addr_line1,
                    addr_line2,
                    city,
                    state,
                    pincode,
                    full_address_str,
                    prod_name,
                    sku,
                    qty,
                    unit_price,
                    item_total,
                    weight,
                    handle_shape,
                    handle_size,
                    grip_color,
                    grip_pattern,
                    grip_count,
                    sticker_finish,
                    pre_knocking,
                    oiling,
                    face_protection,
                    engraving,
                    specs_summary,
                    subtotal_val,
                    gst_percent_val,
                    gst_amount_val,
                    shipping_val,
                    order.coupon_code or "None",
                    discount_val,
                    grand_total_val,
                    order.tracking_number or "Unassigned",
                    order.shipping_carrier or "Unassigned",
                    order.customer_notes or "",
                    order.admin_notes or ""
                ])

        csv_bytes = output.getvalue().encode("utf-8-sig")
        return BytesIO(csv_bytes), "text/csv; charset=utf-8", "csv"
